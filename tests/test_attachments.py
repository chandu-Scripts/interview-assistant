import pytest

from interview_assistant.app.attachments import MAX_CHARS, AttachmentError, read_attachment


def minimal_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for i, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for off in offsets:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    return bytes(out)


def test_text_and_markdown_are_read_as_is(tmp_path):
    (tmp_path / "resume.txt").write_text("  Data engineer, 6 years.  \n", encoding="utf-8")
    (tmp_path / "prompt.md").write_text("# Be concise", encoding="utf-8")
    assert read_attachment(tmp_path / "resume.txt") == "Data engineer, 6 years."
    assert read_attachment(tmp_path / "prompt.md") == "# Be concise"


def test_docx_text_is_extracted(tmp_path):
    import docx

    doc = docx.Document()
    doc.add_paragraph("Senior Azure data engineer")
    doc.add_paragraph("Built streaming pipelines with Databricks")
    path = tmp_path / "resume.docx"
    doc.save(str(path))
    text = read_attachment(path)
    assert "Senior Azure data engineer" in text and "Databricks" in text


def test_pdf_text_is_extracted(tmp_path):
    path = tmp_path / "resume.pdf"
    path.write_bytes(minimal_pdf("Azure Data Engineer"))
    assert "Azure Data Engineer" in read_attachment(path)


def test_unsupported_empty_and_broken_files_are_refused_with_a_reason(tmp_path):
    (tmp_path / "photo.png").write_bytes(b"\x89PNG")
    with pytest.raises(AttachmentError, match="Unsupported file type .png"):
        read_attachment(tmp_path / "photo.png")
    (tmp_path / "blank.txt").write_text("   \n", encoding="utf-8")
    with pytest.raises(AttachmentError, match="no readable text"):
        read_attachment(tmp_path / "blank.txt")
    (tmp_path / "broken.pdf").write_bytes(b"not a pdf")
    with pytest.raises(AttachmentError, match="Could not read broken.pdf"):
        read_attachment(tmp_path / "broken.pdf")


def test_long_attachments_are_capped(tmp_path):
    path = tmp_path / "long.txt"
    path.write_text("x" * (MAX_CHARS + 500), encoding="utf-8")
    assert len(read_attachment(path)) == MAX_CHARS
